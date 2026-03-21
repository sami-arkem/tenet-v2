import csv
import json
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

BASE = Path("data/sources/sanctions")
OUT = Path("data/processed/sanctions/parsed")
OUT.mkdir(parents=True, exist_ok=True)


def write_jsonl(records, out_path):
    count = 0
    with out_path.open("w", encoding="utf-8") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
            count += 1
    return count


def parse_ofac():
    src = BASE / "ofac" / "ofac_sdn.csv"
    out = OUT / "ofac_sdn.jsonl"
    records = []

    with src.open("r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            if not row:
                continue

            records.append(
                {
                    "source": "OFAC",
                    "record_type": "sanctions",
                    "entity_id": row[0].strip() if len(row) > 0 else "",
                    "name": row[1].strip() if len(row) > 1 else "",
                    "program": row[3].strip() if len(row) > 3 else "",
                }
            )

    print(f"ofac rows: {write_jsonl(records, out)}")


def parse_uk():
    src = BASE / "uk" / "uk_sanctions_list.csv"
    out = OUT / "uk_sanctions.jsonl"
    records = []

    with src.open("r", encoding="utf-8", errors="ignore") as f:
        reader = csv.DictReader(f)
        for row in reader:
            records.append(
                {
                    "source": "UK",
                    "record_type": "sanctions",
                    "entity_id": row.get("Unique ID", "").strip(),
                    "name": row.get("Name 6", "").strip() or row.get("Name", "").strip(),
                    "program": row.get("Regime Name", "").strip(),
                }
            )

    print(f"uk rows: {write_jsonl(records, out)}")


def parse_un():
    src = BASE / "un" / "un_consolidated.xml"
    out = OUT / "un_sanctions.jsonl"
    records = []

    tree = ET.parse(src)
    root = tree.getroot()

    for elem in root.iter():
        tag = elem.tag.lower()
        if tag.endswith("individual") or tag.endswith("entity"):
            name_parts = []
            data = {
                "source": "UN",
                "record_type": "sanctions",
                "entity_id": "",
                "name": "",
                "program": "",
            }

            for child in elem.iter():
                ctag = child.tag.lower()
                text = (child.text or "").strip()

                if ctag.endswith("reference_number") and text and not data["entity_id"]:
                    data["entity_id"] = text
                if ctag.endswith("listed_on") and text and not data["program"]:
                    data["program"] = text
                if ctag.endswith("first_name") and text:
                    name_parts.append(text)
                if ctag.endswith("second_name") and text:
                    name_parts.append(text)
                if ctag.endswith("third_name") and text:
                    name_parts.append(text)
                if ctag.endswith("fourth_name") and text:
                    name_parts.append(text)
                if ctag.endswith("name") and text and not name_parts:
                    name_parts.append(text)

            data["name"] = " ".join(part for part in name_parts if part).strip()
            if data["name"] or data["entity_id"]:
                records.append(data)

    print(f"un rows: {write_jsonl(records, out)}")


def parse_eu():
    src = BASE / "eu" / "eu_sanctions.xml"
    out = OUT / "eu_sanctions.jsonl"
    records = []

    tree = ET.parse(src)
    root = tree.getroot()

    for elem in root.iter():
        tag = elem.tag.lower()
        if tag.endswith("sanctionentity"):
            data = {
                "source": "EU",
                "record_type": "sanctions",
                "entity_id": "",
                "name": "",
                "program": "",
            }
            name_parts = []

            for child in elem.iter():
                ctag = child.tag.lower()
                text = (child.text or "").strip()

                if ctag.endswith("eureferencenumber") and text and not data["entity_id"]:
                    data["entity_id"] = text
                if ctag.endswith("regulation") and text and not data["program"]:
                    data["program"] = text
                if ctag.endswith("wholename") and text:
                    name_parts.append(text)

                whole_name_attr = child.attrib.get("wholeName")
                if ctag.endswith("namealias") and whole_name_attr:
                    name_parts.append(whole_name_attr.strip())

            data["name"] = " | ".join(dict.fromkeys([p for p in name_parts if p])).strip()
            if data["name"] or data["entity_id"]:
                records.append(data)

    print(f"eu rows: {write_jsonl(records, out)}")


def parse_canada():
    src = BASE / "canada" / "canada_sanctions.xml"
    out = OUT / "canada_sanctions.jsonl"
    records = []

    tree = ET.parse(src)
    root = tree.getroot()

    for elem in root.iter():
        tag = elem.tag.lower()
        if tag.endswith("record"):
            data = {
                "source": "CANADA",
                "record_type": "sanctions",
                "entity_id": "",
                "name": "",
                "program": "",
            }
            name_parts = []

            for child in elem.iter():
                ctag = child.tag.lower()
                text = (child.text or "").strip()

                if "id" in ctag and text and not data["entity_id"]:
                    data["entity_id"] = text
                if ("country" in ctag or "regulation" in ctag) and text and not data["program"]:
                    data["program"] = text
                if "name" in ctag and text:
                    name_parts.append(text)

            data["name"] = " | ".join(dict.fromkeys([p for p in name_parts if p])).strip()
            if data["name"] or data["entity_id"]:
                records.append(data)

    print(f"canada rows: {write_jsonl(records, out)}")


def parse_australia():
    src = BASE / "australia" / "australia_sanctions.xlsx"
    out = OUT / "australia_sanctions.jsonl"
    records = []

    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        print("australia rows: 0")
        return

    headers = [str(x).strip() if x is not None else "" for x in rows[0]]

    for row in rows[1:]:
        item = dict(zip(headers, row))
        name = str(item.get("Name of Individual or Entity", "") or "").strip()
        ref = str(item.get("Reference", "") or "").strip()
        program = str(item.get("Sanction Imposed", "") or "").strip()

        if name or ref:
            records.append(
                {
                    "source": "AUSTRALIA",
                    "record_type": "sanctions",
                    "entity_id": ref,
                    "name": name,
                    "program": program,
                }
            )

    print(f"australia rows: {write_jsonl(records, out)}")



def parse_new_zealand():
    src = BASE / "new_zealand" / "new_zealand_russia_sanctions.xlsx"
    out = OUT / "new_zealand_sanctions.jsonl"
    records = []

    wb = load_workbook(src, read_only=True, data_only=True)

    for ws in wb.worksheets:
        if ws.title != "Ships":
            continue

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [str(x).strip() if x is not None else "" for x in rows[0]]

        for row in rows[1:]:
            item = dict(zip(headers, row))
            name = str(item.get("Name of Ship as of Date of Sanction", "") or "").strip()
            ref = str(item.get("Unique Identifier", "") or "").strip()
            program = "Ship Ban"

            if name or ref:
                records.append({
                    "source": "NEW_ZEALAND",
                    "record_type": "sanctions",
                    "entity_id": ref,
                    "name": name,
                    "program": program,
                })

    print(f"new_zealand rows: {write_jsonl(records, out)}")


if __name__ == "__main__":
    parse_ofac()
    parse_uk()
    parse_un()
    parse_eu()
    parse_canada()
    parse_australia()
    parse_new_zealand()
